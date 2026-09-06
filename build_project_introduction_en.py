"""Build the English edition of the project introduction."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_project_introduction import (
    BLUE, DARK_BLUE, LIGHT_BLUE, LIGHT_GRAY, MUTED, add_bullet, add_heading,
    add_page_number, add_table, add_text, set_font, set_table_geometry, shade,
)


OUT = Path(__file__).with_name("AI_Multimodal_Detection_System_Project_Introduction_EN.docx")


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("AI Multimodal Detection System | Project Introduction"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("Page "), size=9, color=MUTED)
    add_page_number(footer)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("AI Multimodal Detection System"), size=24, color="0B2545", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("Project Introduction and Technical Roadmap for YOLOE-11 Open-Vocabulary Detection"), size=14, color=MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2700, 6660])
    metadata = [("Document version", "v1.0"), ("Last updated", "31 August 2026"), ("Default model", "YOLOE-11l-seg (open vocabulary, high accuracy)")]
    for row, (label, value) in zip(meta.rows, metadata):
        shade(row.cells[0], LIGHT_GRAY)
        set_font(row.cells[0].paragraphs[0].add_run(label), size=10.5, color=DARK_BLUE, bold=True)
        set_font(row.cells[1].paragraphs[0].add_run(value), size=10.5)

    add_heading(doc, "1. Project Overview")
    add_text(doc, "This project is a Python desktop AI vision application for real-time object detection from a camera or local video. Its core capability is open-vocabulary detection: users can enter natural-language target phrases at runtime, such as 'milk bottle', 'red backpack', or 'mobile phone'. The system uses a vision-language model to locate zero-shot targets instead of being limited to a fixed class list in a conventional detector.")
    add_text(doc, "The project is designed for usability and maintainability. The UI, model inference, video worker, and common utilities are separated into dedicated modules, while CUDA inference, error diagnostics, low-light enhancement, and visual prompting remain available as features or extension points.")

    add_heading(doc, "2. Implemented Capabilities")
    add_table(doc, ["Module", "Current capability", "Verification status"], [
        ("Open text prompts", "Accepts multiple prompts separated by commas, semicolons, or line breaks, including phrase-level object descriptions.", "A real local-image inference was completed with the 'person' prompt on YOLOE-11l."),
        ("High-accuracy model", "Uses YOLOE-11l-seg by default, with YOLOE-11s / 11m / 11l selectable.", "The 11l weights were loaded successfully in a CUDA environment."),
        ("GPU acceleration", "Supports CPU or CUDA and checks CUDA availability during startup.", "Verified on an RTX 5070 Laptop GPU where torch.cuda.is_available() returned true."),
        ("Video input", "Supports Camera 0 and MP4/AVI/MOV/MKV files; DirectShow is preferred for Windows cameras.", "UI and acquisition/error-handling code are complete; continuous streams were not acceptance-tested across every source."),
        ("Black-screen and error handling", "Displays the first camera frame before model warm-up and shows inference exceptions in a dialog.", "Validated through error diagnosis and UI construction checks."),
        ("Visual prompting", "Lets users draw a target box on a reference image and applies HSV color-histogram similarity as a second-stage filter.", "UI and algorithm flow are implemented; no real-scene end-to-end benchmark has yet been completed."),
        ("Low-light enhancement", "Loads Zero-DCE ONNX when available and falls back to CLAHE when it is absent.", "The CLAHE fallback is implemented; the Zero-DCE weight branch awaits validation with a supplied model."),
    ], [2100, 4300, 2960])

    add_heading(doc, "3. Software Structure and Runtime Flow")
    add_table(doc, ["File", "Responsibility"], [
        ("new.py", "Application entry point and controller for model loading, source selection, detection lifecycle, and user-facing errors."),
        ("ui.py", "PySide6 layout construction and the reference-image selection canvas."),
        ("engines.py", "YOLOE-11 open-vocabulary detector plus the Zero-DCE/CLAHE low-light enhancement engine."),
        ("video_worker.py", "QThread-based video acquisition, preview, inference, and status signals."),
        ("app_utils.py", "Reusable helpers, including prompt parsing."),
        ("run_app.bat", "Launches the application with the verified pytorch Conda environment."),
    ], [2500, 6860])
    add_bullet(doc, "Launch: use run_app.bat, or activate the pytorch Conda environment and run python new.py.")
    add_bullet(doc, "Load: select a model scale and device, then click 'Apply Settings / Download and Load YOLOE-11 Model'.")
    add_bullet(doc, "Detect: enter one or more prompts and start the camera or video stream. The first frame appears before model warm-up; detection boxes and confidences follow.")

    add_heading(doc, "4. Key Technologies")
    add_table(doc, ["Technology", "Role in the project"], [
        ("YOLOE-11", "An open-vocabulary detection and instance-segmentation model. It extends conventional YOLO beyond fixed categories through text, visual, and prompt-free mechanisms; this project uses the text-prompt path."),
        ("CLIP / vision-language alignment", "Encodes a user's natural-language query into text embeddings so runtime concepts such as 'milk bottle' can participate in detection."),
        ("PyTorch + CUDA", "Loads YOLOE weights and runs inference on NVIDIA GPUs."),
        ("Ultralytics", "Provides model management, weight loading, prediction, and post-processing interfaces for YOLOE."),
        ("OpenCV", "Handles cameras and video files, color conversion, box drawing, HSV histogram comparison, and CLAHE enhancement."),
        ("PySide6 / Qt", "Builds the desktop GUI and isolates video/inference work from the UI thread with QThread."),
        ("ONNX Runtime", "Provides the optional inference backend for the Zero-DCE low-light enhancement model."),
    ], [2500, 6860])

    add_heading(doc, "5. Current Boundaries and Usage Guidance")
    add_bullet(doc, "Open vocabulary does not mean every concept will be detected reliably. Clear images, common visible objects, and specific correctly spelled English prompts generally produce the best results.")
    add_bullet(doc, "The first use of text prompts requires the official CLIP text-encoder dependency. The project has fixed the prior error caused by installing an unrelated clipboard-manager package with the same name, clip.")
    add_bullet(doc, "YOLOE-11l prioritizes accuracy and therefore requires more VRAM and inference time. Choose YOLOE-11s when real-time responsiveness is more important.")
    add_bullet(doc, "The current visual-prompt path uses color-histogram filtering and is sensitive to illumination, background, and material changes. It should not be treated as deep visual retrieval.")

    add_heading(doc, "6. Future Optimization Directions")
    add_table(doc, ["Direction", "Recommended work", "Expected value"], [
        ("Visual-prompt upgrade", "Adopt YOLOE native visual prompting (SAVPE) or CLIP/DINO feature retrieval instead of relying solely on HSV histograms.", "More robust matching under illumination changes, occlusion, and specific brand/part targets."),
        ("Prompt-free discovery", "Integrate YOLOE prompt-free weights and its built-in vocabulary to propose classes when the user does not know what is present.", "Extends the workflow from finding a named target to exploring a scene."),
        ("Video tracking", "Use ByteTrack or BoT-SORT to assign stable IDs to detections.", "Reduces frame-to-frame jitter and enables counting, trajectories, and alerts."),
        ("Auto-labeling", "Export frames, boxes, and labels to YOLO / COCO / VOC datasets with a human review step.", "Accelerates data collection and domain adaptation."),
        ("Domain fine-tuning", "Collect task-specific data and fine-tune YOLOE / YOLO11 for detection or segmentation.", "Improves accuracy for small targets, unusual angles, and industrial parts."),
        ("Performance engineering", "Add frame skipping, asynchronous queues, mixed precision, TensorRT/ONNX export, and multi-camera resource management.", "Reduces latency and increases throughput."),
        ("Evaluation and explainability", "Add inference logs, FPS/VRAM panels, dataset evaluation, confusion analysis, and replay.", "Provides measurable evidence for model selection and iteration."),
    ], [1900, 4600, 2860])

    add_heading(doc, "7. Recommended Next-Stage Acceptance Tests")
    add_bullet(doc, "Camera validation: run sustained CPU and CUDA sessions and record startup time, FPS, VRAM use, and error recovery behavior.")
    add_bullet(doc, "Prompt-set evaluation: prepare at least 30 target phrases across people, containers, tools, electronics, and small objects.")
    add_bullet(doc, "Video and low-light tests: compare raw video, CLAHE, and Zero-DCE on local files and practical low-light scenes.")
    add_bullet(doc, "Visual-prompt acceptance: prepare samples across lighting and backgrounds, then compare recall and false positives before and after the planned upgrade.")

    add_heading(doc, "8. Reference Papers and Resources")
    references = [
        "[1] Wang, A. et al. YOLOE: Real-Time Seeing Anything. arXiv:2503.07465, 2025. https://arxiv.org/abs/2503.07465",
        "[2] Radford, A. et al. Learning Transferable Visual Models from Natural Language Supervision. ICML, 2021, pp. 8748-8763. https://mlanthology.org/icml/2021/radford2021icml-learning/",
        "[3] Cheng, T. et al. YOLO-World: Real-Time Open-Vocabulary Object Detection. CVPR, 2024. https://arxiv.org/abs/2401.17270",
        "[4] Guo, C. et al. Zero-Reference Deep Curve Estimation for Low-Light Image Enhancement. CVPR, 2020, pp. 1780-1789. https://openaccess.thecvf.com/content_CVPR_2020/html/Guo_Zero-Reference_Deep_Curve_Estimation_for_Low-Light_Image_Enhancement_CVPR_2020_paper.html",
        "[5] Ultralytics Documentation. YOLOE Open-Vocabulary Detection & Segmentation. https://docs.ultralytics.com/models/yoloe/",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.10
        set_font(p.add_run(reference), size=9.5)

    doc.core_properties.title = "AI Multimodal Detection System - Project Introduction"
    doc.core_properties.subject = "YOLOE-11 open-vocabulary detection project overview"
    doc.core_properties.author = "AI Multimodal Detection System Project"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
