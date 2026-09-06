from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("AI多模态检测系统_项目介绍.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def set_cell_width(cell, width_dxa):
    props = cell._tc.get_or_add_tcPr()
    width = props.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        props.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_props = table._tbl.tblPr
    table_width = table_props.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_props.append(table_width)
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_props.append(indent)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")
                mar.append(node)
            margins.append(mar)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), bold=True)
        set_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    set_font(p.add_run(text))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 13, color=BLUE, bold=True)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=10.5, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("AI 多模态检测系统 | 项目介绍"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("第 "), size=9, color=MUTED)
    add_page_number(footer)
    set_font(footer.add_run(" 页"), size=9, color=MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("AI 多模态检测系统"), size=24, color="0B2545", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("YOLOE-11 开放词汇目标检测项目介绍与技术路线"), size=14, color=MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2700, 6660])
    metadata = [("文档版本", "v1.0"), ("更新日期", "2026 年 8 月 31 日"), ("当前默认模型", "YOLOE-11l-seg（开放词汇、高精度）")]
    for row, (label, value) in zip(meta.rows, metadata):
        shade(row.cells[0], LIGHT_GRAY)
        set_font(row.cells[0].paragraphs[0].add_run(label), size=10.5, color=DARK_BLUE, bold=True)
        set_font(row.cells[1].paragraphs[0].add_run(value), size=10.5)

    add_heading(doc, "1. 项目概述")
    add_text(doc, "本项目是一个基于 Python 的桌面端 AI 视觉应用，用于在摄像头或本地视频中进行实时目标检测。当前核心能力是开放词汇检测：用户可在运行时输入自然语言目标短语，例如 “milk bottle”“red backpack” 或 “mobile phone”，系统利用视觉-语言模型完成零样本目标定位，而不是受限于传统检测模型的固定类别表。")
    add_text(doc, "项目以可用性和可维护性为目标：将界面、模型推理、视频线程与通用工具拆分为独立模块，同时保留 CUDA 推理、错误诊断、低照度增强与视觉提示词等功能入口。")

    add_heading(doc, "2. 当前已实现功能")
    add_table(doc, ["模块", "当前能力", "验证状态"], [
        ("开放文本提示词", "接受逗号、分号或换行分隔的多个提示词；支持短语级目标描述。", "已在 YOLOE-11l 上用 person 提示词和本地图片完成真实推理。"),
        ("高精度模型", "默认使用 YOLOE-11l-seg；可切换 YOLOE-11s / 11m / 11l。", "已在 CUDA 环境成功加载 11l 权重。"),
        ("GPU 加速", "CPU / CUDA 可选；启动时检查 CUDA 是否可用。", "已确认 RTX 5070 Laptop GPU 环境中 torch.cuda.is_available() 为真。"),
        ("视频输入", "支持 Camera 0 和 MP4/AVI/MOV/MKV 本地文件；Windows 摄像头优先 DirectShow。", "UI 与采集/错误处理代码已完成；持续视频流未作为发布验收项逐源测试。"),
        ("黑屏与异常诊断", "首帧在模型预热前显示；推理异常弹窗显示具体异常。", "已通过异常定位与 UI 构建验证。"),
        ("视觉提示词", "参考图拖拽框选目标后，使用 HSV 颜色直方图做二次相似度筛选。", "界面与算法流程已实现；尚未做真实场景端到端评测。"),
        ("暗光增强", "优先加载 Zero-DCE ONNX；缺失时退化为 CLAHE。", "CLAHE 回退逻辑已实现；Zero-DCE 权重分支待提供模型后验证。"),
    ], [2100, 4300, 2960])

    add_heading(doc, "3. 软件结构与运行流程")
    add_table(doc, ["文件", "职责"], [
        ("new.py", "应用入口与控制器：模型加载、输入源选择、启动/停止检测、错误提示。"),
        ("ui.py", "PySide6 界面构建与参考图片框选画布。"),
        ("engines.py", "YOLOE-11 开放词汇检测引擎与 Zero-DCE/CLAHE 暗光增强引擎。"),
        ("video_worker.py", "QThread 视频采集、预览、推理与状态信号。"),
        ("app_utils.py", "提示词解析等可复用工具。"),
        ("run_app.bat", "使用已验证的 pytorch Conda 环境启动程序。"),
    ], [2500, 6860])
    add_bullet(doc, "启动：使用 run_app.bat 或在 pytorch Conda 环境中运行 python new.py。")
    add_bullet(doc, "加载：选择模型规模与设备，点击“应用设置 / 下载并加载 YOLOE-11 模型”。")
    add_bullet(doc, "检测：输入一个或多个文本提示词，启动摄像头/视频流；首帧先显示，随后叠加检测框和置信度。")

    add_heading(doc, "4. 关键技术")
    add_table(doc, ["技术", "在项目中的作用"], [
        ("YOLOE-11", "开放词汇检测与实例分割模型。通过文本、视觉或无提示词机制扩展传统 YOLO 的固定类别限制；本项目使用文本提示词检测路径。"),
        ("CLIP / 视觉-语言对齐", "将用户输入的自然语言转为文本嵌入，使 “milk bottle” 等运行时类别可参与检测。"),
        ("PyTorch + CUDA", "加载 YOLOE 权重并调用 NVIDIA GPU 加速推理。"),
        ("Ultralytics", "提供 YOLOE 模型管理、权重加载、预测与后处理接口。"),
        ("OpenCV", "摄像头与视频读取、颜色空间转换、绘制框、HSV 直方图相似度和 CLAHE 图像增强。"),
        ("PySide6 / Qt", "构建桌面 GUI，并以 QThread 将视频/推理从界面线程中分离。"),
        ("ONNX Runtime", "为可选的 Zero-DCE ONNX 暗光增强模型提供推理后端。"),
    ], [2500, 6860])

    add_heading(doc, "5. 已知边界与使用建议")
    add_bullet(doc, "“开放词汇”并不等于任何概念都能稳定识别。清晰画面、常见可见物体、具体且拼写正确的英文短语通常效果更好。")
    add_bullet(doc, "文本提示词的首次使用需要官方 CLIP 文本编码依赖；项目已修复误装同名 clipboard clip 包导致的属性错误。")
    add_bullet(doc, "YOLOE-11l 追求精度，需要更多显存与推理时间；实时性优先时应选择 YOLOE-11s。")
    add_bullet(doc, "当前视觉提示词使用颜色直方图二次筛选，对光照、背景和材质变化敏感，不应视为深度视觉检索能力。")

    add_heading(doc, "6. 后续优化方向")
    add_table(doc, ["方向", "建议工作", "预期价值"], [
        ("视觉提示词升级", "接入 YOLOE 原生视觉提示词（SAVPE）或 CLIP/DINO 特征检索，替换单纯 HSV 直方图。", "提升对颜色变化、遮挡和特定品牌/部件目标的稳健性。"),
        ("无提示词发现", "接入 YOLOE prompt-free 权重与内置大词表，为未知画面提供候选类别。", "从“找指定目标”扩展为“探索画面中有什么”。"),
        ("视频追踪", "使用 ByteTrack 或 BoT-SORT 为检测目标分配稳定 ID。", "减少逐帧框抖动，支持计数、轨迹和告警。"),
        ("自动标注", "将视频帧、检测框和类别导出为 YOLO / COCO / VOC 数据集格式，并加入人工复核。", "加速特定业务场景的数据积累与微调。"),
        ("领域微调", "围绕实际对象采集数据，使用 YOLOE / YOLO11 进行检测或分割微调。", "提高小目标、特殊角度、工业零件等领域场景精度。"),
        ("性能工程", "加入帧跳过、异步队列、半精度、TensorRT/ONNX 导出与多摄像头资源管理。", "降低延迟并提高吞吐量。"),
        ("可解释与评估", "增加推理日志、FPS/显存面板、数据集评测、混淆分析与回放。", "使模型选择与迭代有量化依据。"),
    ], [1900, 4600, 2860])

    add_heading(doc, "7. 建议的下一阶段验收")
    add_bullet(doc, "摄像头实测：分别在 CPU 与 CUDA 下连续运行，记录启动耗时、FPS、显存占用与异常恢复情况。")
    add_bullet(doc, "提示词集评测：建立不少于 30 个目标短语的测试表，覆盖人物、容器、工具、电子产品和小目标。")
    add_bullet(doc, "视频与弱光测试：用本地视频和实际低照度场景比较原始画面、CLAHE 与 Zero-DCE 的检测效果。")
    add_bullet(doc, "视觉提示词验收：准备多光照、多背景样本，确认升级前后相似目标检索的召回率和误检率。")

    add_heading(doc, "8. 参考论文与资料")
    refs = [
        "[1] Wang, A. et al. YOLOE: Real-Time Seeing Anything. arXiv:2503.07465, 2025. https://arxiv.org/abs/2503.07465",
        "[2] Radford, A. et al. Learning Transferable Visual Models from Natural Language Supervision. ICML, 2021, pp. 8748-8763. https://mlanthology.org/icml/2021/radford2021icml-learning/",
        "[3] Cheng, T. et al. YOLO-World: Real-Time Open-Vocabulary Object Detection. CVPR, 2024. https://arxiv.org/abs/2401.17270",
        "[4] Guo, C. et al. Zero-Reference Deep Curve Estimation for Low-Light Image Enhancement. CVPR, 2020, pp. 1780-1789. https://openaccess.thecvf.com/content_CVPR_2020/html/Guo_Zero-Reference_Deep_Curve_Estimation_for_Low-Light_Image_Enhancement_CVPR_2020_paper.html",
        "[5] Ultralytics Documentation. YOLOE Open-Vocabulary Detection & Segmentation. https://docs.ultralytics.com/models/yoloe/",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.0)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.10
        set_font(p.add_run(ref), size=9.5)

    doc.core_properties.title = "AI 多模态检测系统 - 项目介绍"
    doc.core_properties.subject = "YOLOE-11 开放词汇检测项目技术说明"
    doc.core_properties.author = "AI 多模态检测系统项目"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
